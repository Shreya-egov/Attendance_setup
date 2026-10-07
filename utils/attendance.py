"""
Attendance planning: which boundaries get a register, and which user goes on
which register as attendee / OWNER / APPROVER. Pure functions over search
results - the writes are Hcm.create_registers / create_staff /
create_attendees. The reasoning behind the direct API route is in
docs/register-attendee-api-chain.md.
"""
import datetime as dt

from utils.client import die, epoch_ms, fmt_date
from utils.config import (
    AUTO_LEVEL_ROLE, MULTI_REGISTER_ROLES, TENANT,
)


def resolve_enrollment(campaign, arg):
    """
    Enrollment date, defaulting to campaign start. The direct API route does
    not run AttendanceRegisterAttendeeValidationProcessor (confirmed on demo
    2026-09-28: that processor only fires on the xlsx route), so the
    out-of-range rejection that guards the sheet route is absent here and we
    enforce the window ourselves rather than letting bad data through.
    """
    if arg:
        try:
            d = dt.datetime.strptime(arg, "%d-%m-%Y").date()
        except ValueError:
            die(f"--enrollment-date must be DD-MM-YYYY, got {arg!r}")
        ms = epoch_ms(d)
    else:
        ms = campaign.get("startDate")

    start, end = campaign.get("startDate"), campaign.get("endDate")
    if ms and start and ms < start:
        die(f"enrollment date {fmt_date(ms)} is before the campaign start "
            f"{fmt_date(start)}.\n"
            "         That must be rejected, not clamped - refusing to send it.")
    if ms and end and ms > end:
        die(f"enrollment date {fmt_date(ms)} is after the campaign end {fmt_date(end)}")
    return ms


# ---- planning ----------------------------------------------------------


def boundary_levels(campaign):
    """code -> hierarchy type (COUNTRY, PROVINCE, DISTRICT, ... PHU)."""
    return {b["code"]: b.get("type") for b in campaign.get("boundaries") or []
            if b.get("code")}


def select_boundaries(hcm, campaign, bp_map, level, all_levels=False, users=None):
    """
    Narrow boundary -> projectId to the level registers go at (see
    resolve_level; "leaf" = deepest; "auto" = where AUTO_LEVEL_ROLE users
    are mapped, needs users; all_levels keeps everything). Boundary types
    come from the campaign's list plus the projects' own addresses.
    """
    levels = {**boundary_levels(campaign), **hcm.boundary_types}
    if all_levels:
        return bp_map
    ordered = hcm.hierarchy_levels(campaign["hierarchyType"])
    if str(level).lower() == "auto":
        return auto_level_boundaries(bp_map, levels, ordered, users or [])
    if str(level).lower() == "leaf":
        leaves = leaf_boundaries(campaign, ordered)
        kept = {k: v for k, v in bp_map.items() if k in leaves}
        if len(kept) < len(bp_map):
            print(f"           {len(bp_map) - len(kept)} non-leaf boundary/ies skipped"
                  f" (--all-levels to include)")
        return kept
    want = resolve_level(level, ordered)
    print(f"           creating at level {ordered.index(want) + 1} = {want}"
          f"  (hierarchy: {' > '.join(ordered)})")
    kept = {k: v for k, v in bp_map.items() if (levels.get(k) or "").upper() == want}
    if not kept:
        die(f"no campaign boundary of level {want} has a project. Available: "
            f"{sorted({v for v in levels.values() if v})}")
    return kept


def auto_level_boundaries(bp_map, levels, ordered, users, role=None):
    """
    Registers go where the distributors are: the boundaries AUTO_LEVEL_ROLE
    users are mapped to in their credentials, at the level most of them sit
    at. Boundaries at that level with no distributor get no register - it
    would stay empty.
    """
    role = role or AUTO_LEVEL_ROLE
    dist = [u for u in users if role in u["roles"]]
    if not dist:
        die(f"--level auto: no {role} user in this campaign to take the level from.\n"
            f"         Pass --level N or a type ({', '.join(ordered)}).")
    by_type = {}
    for u in dist:
        by_type.setdefault((levels.get(u["boundaryCode"]) or "?").upper(), set()).add(
            u["boundaryCode"])
    want = max(by_type, key=lambda t: len(by_type[t]))
    if want == "?":
        die(f"--level auto: cannot tell the boundary type of the {role} users' "
            f"boundaries {sorted(by_type['?'])[:5]}; pass --level")
    if len(by_type) > 1:
        print(f"  WARNING  {role} users sit at several levels "
              f"{ {t: len(c) for t, c in by_type.items()} }; using {want}")
    print(f"           creating at level {ordered.index(want) + 1} = {want}"
          f"  (auto: where {len(dist)} {role} user(s) are mapped)")
    kept = {c: bp_map[c] for c in by_type[want] if c in bp_map}
    no_project = sorted(by_type[want] - set(bp_map))
    if no_project:
        print(f"  WARNING  {len(no_project)} {role} boundary/ies have no project and get"
              f" no register: {no_project[:5]}")
    if not kept:
        die(f"--level auto: none of the {role} boundaries has a project")
    return kept


def resolve_level(arg, ordered):
    """
    --level value -> boundary type. Accepts a 1-based position from the root
    ("3" -> DISTRICT in NIGERIA) or a type name ("district").
    """
    arg = str(arg).strip()
    shown = ", ".join(f"{i}={t}" for i, t in enumerate(ordered, 1))
    if arg.isdigit():
        n = int(arg)
        if not 1 <= n <= len(ordered):
            die(f"--level {n} is out of range. Hierarchy levels: {shown}")
        return ordered[n - 1]
    want = arg.upper()
    if want not in ordered:
        die(f"--level {arg!r} is not a type in this hierarchy. Levels: {shown}")
    return want


def leaf_boundaries(campaign, ordered=None):
    """
    Codes of the campaign boundaries that have no child in the campaign's own
    boundary set. The project service returns a project for every hierarchy
    level (COUNTRY -> PROVINCE -> ... -> PHU), but a register belongs at the
    operational level, so by default only leaves get one.

    unified-qa campaigns carry no "parent" on their boundaries, which would
    make every boundary a leaf. Without parents, fall back to the boundaries
    at the deepest hierarchy level present (needs ordered, root-first types).
    """
    bs = campaign.get("boundaries") or []
    codes = {b.get("code") for b in bs if b.get("code")}
    parents = {b.get("parent") for b in bs if b.get("parent")}
    if parents or not ordered:
        return {c for c in codes if c not in parents}
    present = [t for t in ordered if any(b.get("type") == t for b in bs)]
    deepest = present[-1] if present else None
    return {b["code"] for b in bs if b.get("code") and b.get("type") == deepest}


def boundary_project_map(rows):
    """
    boundaryCode -> projectId, exactly as project-factory builds it
    (attendanceRegister-processClass.ts buildBoundaryProjectMap): boundary rows
    carry the code in data and the created projectId in uniqueIdAfterProcess.
    """
    out = {}
    for row in rows:
        data = row.get("data") or {}
        code = data.get("HCM_ADMIN_CONSOLE_BOUNDARY_CODE") or row.get("uniqueIdentifier")
        project_id = row.get("uniqueIdAfterProcess")
        if code and project_id:
            out[code] = project_id
    return out


def campaign_users(rows):
    """
    user rows -> [{userId, boundaryCode, roles[], username}]. uniqueIdAfterProcess
    is the HRMS user uuid, which is what the attendance APIs want for both
    attendee.individualId and staff.userId.
    """
    users = []
    for row in rows:
        data = row.get("data") or {}
        user_id = row.get("uniqueIdAfterProcess")
        if not user_id:
            continue
        roles_raw = data.get("HCM_ADMIN_CONSOLE_USER_ROLE") or ""
        roles = {
            r.strip().upper().replace(" ", "_")
            for r in str(roles_raw).replace(",", "#").split("#")
            if r.strip()
        }
        users.append({
            "userId": user_id,
            "username": data.get("UserName") or row.get("uniqueIdentifier"),
            "boundaryCode": data.get("HCM_ADMIN_CONSOLE_BOUNDARY_CODE_MANDATORY")
                            or data.get("HCM_ADMIN_CONSOLE_BOUNDARY_CODE"),
            "roles": roles,
        })
    return users


def plan_registers(campaign, bp_map, existing, prefix, event_type, sessions,
                   force=False, name_suffix=""):
    """
    One register per campaign boundary that has a project and no register yet.

    force=True skips the already-covered check, so a boundary that already has
    a register gets another one. The API does not enforce uniqueness itself;
    the guard is ours.
    """
    by_locality = set() if force else {r.get("localityCode") for r in existing}
    name = campaign.get("campaignName") or campaign.get("campaignNumber")
    number = campaign["campaignNumber"]
    start, end = campaign.get("startDate"), campaign.get("endDate")
    seq = len(existing)
    planned = []
    for code, project_id in sorted(bp_map.items()):
        if code in by_locality:
            continue
        seq += 1
        planned.append({
            "tenantId": TENANT,
            "name": f"{name} {code}{name_suffix}",
            "referenceId": project_id,
            "campaignNumber": number,
            "serviceCode": f"{prefix}-{number.split('-')[-1]}-{seq:03d}",
            "startDate": start,
            "endDate": end,
            "localityCode": code,
            "additionalDetails": {
                "campaignNumber": number,
                "campaignName": name,
                "eventType": event_type,
                "sessions": sessions,
            },
        })
    return planned


def read_users_xlsx(path):
    """
    Rows of the workbench Users download that were actually created
    (#status# CREATED), as {username, name, userUuid, boundaryCode, roles}.
    Row 2 is the localized header row and has no #status#, so it drops out.
    """
    try:
        import openpyxl
    except ImportError:
        die("reading --users-xlsx needs openpyxl: pip install openpyxl")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as e:
        die(f"cannot open {path}: {e}")
    sheet = "Create List of Users"
    if sheet not in wb.sheetnames:
        die(f"{path} has no {sheet!r} sheet (found {wb.sheetnames})")
    rows = wb[sheet].iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows, [])]
    col = {h: i for i, h in enumerate(header) if h}
    for need in ("UserName", "HCM_ADMIN_CONSOLE_BOUNDARY_CODE_MANDATORY",
                 "UserService Uuids", "#status#"):
        if need not in col:
            die(f"{path}: column {need!r} not found")
    role_cols = [i for h, i in col.items()
                 if h.startswith("HCM_ADMIN_CONSOLE_USER_ROLE")]

    def cell(row, name):
        i = col.get(name)
        v = row[i] if i is not None and i < len(row) else None
        return str(v).strip() if v is not None else ""

    out = []
    for row in rows:
        if cell(row, "#status#").upper() != "CREATED" or not cell(row, "UserService Uuids"):
            continue
        roles = {
            str(row[i]).strip().upper().replace(" ", "_")
            for i in role_cols if i < len(row) and row[i] and str(row[i]).strip()
        }
        out.append({
            "username": cell(row, "UserName"),
            "name": cell(row, "HCM_ADMIN_CONSOLE_USER_NAME") or cell(row, "UserName"),
            "userUuid": cell(row, "UserService Uuids"),
            "boundaryCode": cell(row, "HCM_ADMIN_CONSOLE_BOUNDARY_CODE_MANDATORY"),
            "roles": roles,
        })
    return out


def classify_user(roles, marker_roles, approver_roles, worker_roles):
    """OWNER / APPROVER / ATTENDEE / None. Priority APPROVER > MARKER > WORKER,
    first match wins - the same order excel-ingestion's classifyUserToSheet uses."""
    if roles & approver_roles:
        return "APPROVER"
    if roles & marker_roles:
        return "OWNER"
    if roles & worker_roles:
        return "ATTENDEE"
    return None


def plan_attendance(registers, users, enrollment_ms, marker_roles, approver_roles,
                    worker_roles, multi_roles=MULTI_REGISTER_ROLES, all_registers=None,
                    paths=None):
    """
    Map each user onto in-scope register(s), matched through the boundary
    tree the way the xlsx route's boundaryFilter does (paths: code ->
    [root..code], from Hcm.boundary_paths; without it, exact boundary only):
      workers, markers (LEVEL_RANGE)      the register at the user's boundary
                                          or its nearest ancestor that has one
      approvers (ANCESTOR_AND_SELF)       every register at or below the
                                          user's boundary
    Roles:
      approver roles -> staff APPROVER
      marker roles   -> staff OWNER
      worker roles   -> attendee
      anything else  -> not enrolled (returned in `unmapped`)
    Mirrors the worker/marker/approver sheets of the xlsx route.

    Staff go on every register they match - their roles may sit on several.
    A worker may be an active attendee of only ONE register campaign-wide
    (AttendanceRegisterAttendeeValidationProcessor), so a worker already on
    any register in `all_registers` is skipped, and a worker matching several
    in-scope registers is returned in
    `ambiguous` rather than guessed: narrow with --register-id.

    Staff carry additionalDetails {ownerName, staffName} as the UI writes them,
    ownerName being the register's OWNER (existing or planned).

    Returns a dict of lists: attendees, staff, unplaced, skipped, unmapped,
    ambiguous, and no_owner - registers that will have no OWNER, so nobody
    can mark attendance on them.
    """
    by_locality = {}
    for r in registers:
        by_locality.setdefault(r.get("localityCode"), []).append(r)

    # Already-enrolled sets per register. These APIs are not idempotent - a
    # second _create for the same (register, user) either duplicates or 400s,
    # so skip them here rather than finding out per batch.
    enrolled_attendees, enrolled_staff, owner_name = {}, {}, {}
    for r in registers:
        rid = r["id"]
        enrolled_attendees[rid] = {
            a.get("individualId") for a in r.get("attendees") or []
            if not a.get("denrollmentDate")
        }
        enrolled_staff[rid] = {
            (s.get("userId"), s.get("staffType")) for s in r.get("staff") or []
        }
        for s in r.get("staff") or []:
            if s.get("staffType") == "OWNER":
                owner_name.setdefault(
                    rid, (s.get("additionalDetails") or {}).get("staffName"))
    on_any_register = {
        a.get("individualId")
        for r in (all_registers if all_registers is not None else registers)
        for a in r.get("attendees") or [] if not a.get("denrollmentDate")
    }

    paths = paths or {}

    def path(code):
        return paths.get(code) or [code]

    def registers_for(u, kind):
        if kind == "APPROVER":
            return [r for r in registers if u["boundaryCode"] in path(r.get("localityCode"))]
        for code in reversed(path(u["boundaryCode"])):
            if code in by_locality:
                return by_locality[code]
        return []

    out = {k: [] for k in
           ("attendees", "staff", "unplaced", "skipped", "unmapped", "ambiguous")}
    pending_staff = []
    for u in users:
        kind = classify_user(u["roles"], marker_roles, approver_roles, worker_roles)
        if kind is None:
            out["unmapped"].append(u)
            continue
        regs = registers_for(u, kind)
        if not regs:
            out["unplaced"].append(u)
            continue

        if kind == "ATTENDEE":
            if u["userId"] in on_any_register and not (u["roles"] & multi_roles):
                out["skipped"].append(u)
                continue
            if len(regs) > 1:
                out["ambiguous"].append(u)
                continue
            rid = regs[0]["id"]
            if u["userId"] in enrolled_attendees[rid]:
                out["skipped"].append(u)
                continue
            out["attendees"].append({
                "tenantId": TENANT, "registerId": rid,
                "individualId": u["userId"], "enrollmentDate": enrollment_ms,
            })
            on_any_register.add(u["userId"])
            continue

        for r in regs:
            rid = r["id"]
            if (u["userId"], kind) in enrolled_staff[rid]:
                out["skipped"].append(u)
                continue
            if kind == "OWNER":
                owner_name.setdefault(rid, u["name"])
            pending_staff.append((rid, kind, u))

    has_owner = {rid for rid, kind, _ in pending_staff if kind == "OWNER"} | {
        r["id"] for r in registers
        if any(st.get("staffType") == "OWNER" for st in r.get("staff") or [])}
    out["no_owner"] = [r for r in registers if r["id"] not in has_owner]

    for rid, kind, u in pending_staff:
        out["staff"].append({
            "tenantId": TENANT, "registerId": rid, "userId": u["userId"],
            "staffType": kind, "enrollmentDate": enrollment_ms,
            "additionalDetails": {
                "ownerName": owner_name.get(rid) or u["name"],
                "staffName": u["name"],
            },
        })
    return out


def report_registers(hcm, number):
    """Re-search and print registers with attendee/staff counts."""
    final, body = hcm.registers(number, detailed=True)
    n_att = sum(len(r.get("attendees") or []) for r in final)
    n_staff = sum(len(r.get("staff") or []) for r in final)
    print(f"\nverify     {body.get('totalCount')} register(s),"
          f" {n_att} attendee(s), {n_staff} staff")
    for r in final:
        print(f"           {r['registerNumber']}  {r.get('serviceCode')}"
              f"  {r.get('localityCode')}"
              f"  attendees={len(r.get('attendees') or [])}"
              f"  staff={len(r.get('staff') or [])}")
    return final


# ---- map_staff helpers -------------------------------------------------


def role_set(csv):
    return {r.strip().upper().replace(" ", "_") for r in csv.split(",") if r.strip()}


def unique(users):
    """Users once each - staff are skipped per register, so repeat."""
    seen, out = set(), []
    for u in users:
        if u["userId"] not in seen:
            seen.add(u["userId"])
            out.append(u)
    return out


def names(users, n=6):
    shown = [u.get("username") or u.get("userId", "")[:8] for u in users[:n]]
    return ", ".join(shown) + (" …" if len(users) > n else "")


def warn_no_owner(plan):
    if plan.get("no_owner"):
        regs = plan["no_owner"]
        shown = ", ".join(f"{r.get('serviceCode')} ({r.get('localityCode')})" for r in regs[:5])
        print(f"  WARNING  {len(regs)} register(s) will have no marker (OWNER) - nobody can"
              f" mark attendance on them: {shown}{' …' if len(regs) > 5 else ''}")
